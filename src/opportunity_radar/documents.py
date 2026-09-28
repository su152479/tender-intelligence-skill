"""Attachment discovery, bounded download, text extraction and local indexing."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from hashlib import sha256
from io import BytesIO
import json
import mimetypes
import os
from pathlib import Path
import re
import subprocess
from urllib.parse import unquote, urljoin, urlsplit
import zipfile

from bs4 import BeautifulSoup
import requests

from .collectors.http import HostRateLimiter, PublicPageClient, SourceCircuitOpen
from .config import DATA_DIR
from .db import Database
from .network import NetworkPolicy


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".zip"}
DOCUMENT_LINK_MARKERS = ("附件", "下载", "初步设计", "施工图", "招标文件", "工程量清单")
IDENTIFIER_PATTERNS = (
    ("交易项目编号", re.compile(r"\bS\d{6,}[A-Za-z0-9_-]*\b", re.IGNORECASE)),
    ("工程编号", re.compile(r"\b20\d{2}[-－—][A-Za-z0-9]{2,24}\b", re.IGNORECASE)),
    ("项目编号", re.compile(r"(?:工程|项目|交易|招标)编号[：:\s]*([A-Za-z0-9][A-Za-z0-9_.-]{4,40})", re.IGNORECASE)),
)


@dataclass(frozen=True)
class AttachmentLink:
    url: str
    file_name: str


@dataclass
class ParsedDocument:
    title: str
    text: str
    page_count: int
    parser: str
    metadata: dict


class DocumentParser:
    def __init__(self):
        self.max_text_chars = int(os.getenv("RADAR_DOCUMENT_MAX_TEXT_CHARS", "2000000"))
        self.max_zip_members = int(os.getenv("RADAR_DOCUMENT_MAX_ZIP_MEMBERS", "100"))
        self.max_zip_uncompressed = int(os.getenv("RADAR_DOCUMENT_MAX_ZIP_UNCOMPRESSED_BYTES", str(200 * 1024 * 1024)))

    def parse(self, path: Path) -> ParsedDocument:
        suffix = path.suffix.lower()
        with path.open("rb") as stream:
            return self.parse_bytes(stream.read(), suffix, path.name)

    def parse_bytes(self, payload: bytes, suffix: str, title: str) -> ParsedDocument:
        suffix = suffix.lower()
        if suffix == ".pdf":
            return self._pdf(payload, title)
        if suffix == ".docx":
            return self._docx(payload, title)
        if suffix == ".xlsx":
            return self._xlsx(payload, title)
        if suffix == ".zip":
            return self._zip(payload, title)
        raise ValueError(f"不支持的附件格式：{suffix or '未知'}")

    def _pdf(self, payload: bytes, title: str) -> ParsedDocument:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(payload), strict=False)
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        metadata = {str(key): str(value) for key, value in (reader.metadata or {}).items()}
        metadata["needs_ocr"] = len(text.strip()) < max(200, len(reader.pages) * 20)
        return ParsedDocument(title, self._limit(text), len(reader.pages), "pypdf", metadata)

    def _docx(self, payload: bytes, title: str) -> ParsedDocument:
        from docx import Document

        document = Document(BytesIO(payload))
        values = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            values.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return ParsedDocument(title, self._limit("\n".join(values)), 0, "python-docx", {})

    def _xlsx(self, payload: bytes, title: str) -> ParsedDocument:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        values, cell_count = [], 0
        for sheet in workbook.worksheets[:20]:
            values.append(f"[工作表] {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value) for value in row if value not in (None, "")]
                cell_count += len(cells)
                if cells:
                    values.append("\t".join(cells))
                if cell_count >= 100000:
                    break
            if cell_count >= 100000:
                break
        return ParsedDocument(title, self._limit("\n".join(values)), 0, "openpyxl", {"cells": cell_count})

    def _zip(self, payload: bytes, title: str) -> ParsedDocument:
        values, parsed_members, total_size = [], [], 0
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            members = [item for item in archive.infolist() if not item.is_dir()]
            if len(members) > self.max_zip_members:
                raise ValueError(f"ZIP 文件数超过限制：{len(members)}")
            for item in members:
                if item.flag_bits & 0x1:
                    raise ValueError("ZIP 包含加密文件，已停止解析")
                total_size += item.file_size
                if total_size > self.max_zip_uncompressed:
                    raise ValueError("ZIP 解压后体积超过限制")
                suffix = Path(item.filename).suffix.lower()
                if suffix not in SUPPORTED_EXTENSIONS - {".zip"}:
                    continue
                nested = self.parse_bytes(archive.read(item), suffix, Path(item.filename).name)
                values.append(f"\n===== {item.filename} =====\n{nested.text}")
                parsed_members.append(item.filename)
        return ParsedDocument(
            title, self._limit("\n".join(values)), 0, "zip-container",
            {"parsed_members": parsed_members, "uncompressed_bytes": total_size},
        )

    def _limit(self, value: str) -> str:
        return re.sub(r"\x00", "", value or "")[: self.max_text_chars]


class AttachmentEnricher:
    def __init__(self, db: Database):
        self.db = db
        self.parser = DocumentParser()
        self.data_dir = DATA_DIR / "documents"
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.max_bytes = int(os.getenv("RADAR_DOCUMENT_MAX_BYTES", str(100 * 1024 * 1024)))
        self.minimum_interval = float(os.getenv("RADAR_HTTP_MIN_INTERVAL_SECONDS", "2.5"))
        self.timeout = int(os.getenv("RADAR_COLLECTION_TIMEOUT_SECONDS", "60"))
        self.network = NetworkPolicy.from_env()

    def enrich_project(
        self, project, extra_pages: list[str] | None = None,
        extra_attachments: list[str] | None = None,
    ) -> dict:
        page_urls = list(dict.fromkeys([project["url"], *(extra_pages or [])]))
        links: dict[str, AttachmentLink] = {}
        for page_url in page_urls:
            client = PublicPageClient(f"附件发现:{project['source_site']}", f"documents-{project['id']}")
            html = client.get_detail_text(page_url, allow_windows_curl_fallback=True)
            for link in self.discover_links(html, page_url):
                if self._host_allowed(page_url, link.url):
                    links[link.url] = link
        for url in extra_attachments or []:
            if not self._host_allowed(project["url"], url):
                raise ValueError(f"补充附件不在允许的官方域名内：{url}")
            file_name = Path(unquote(urlsplit(url).path)).name or f"attachment-{sha256(url.encode()).hexdigest()[:12]}"
            links[url] = AttachmentLink(url, sanitize_file_name(file_name))
        result = {
            "notice_id": project["id"], "discovered": len(links),
            "parsed": 0, "skipped": 0, "failed": 0,
        }
        for link in links.values():
            attachment_id = self.db.upsert_attachment(project["id"], link.url, link.file_name)
            existing = self.db.attachment(attachment_id)
            if existing and (
                existing["status"] == "PURGED"
                or (
                    existing["status"] == "PARSED" and existing["local_path"]
                    and Path(existing["local_path"]).is_file()
                )
            ):
                result["skipped"] += 1
                continue
            try:
                path, content_type, size, digest = self.download(link, attachment_id)
                parsed = self.parser.parse(path)
                identifiers = extract_identifiers(parsed.text)
                self.db.update_attachment(
                    attachment_id, content_type=content_type, size_bytes=size, sha256=digest,
                    status="PARSED", local_path=str(path), error="",
                    downloaded_at=datetime.now().isoformat(timespec="seconds"),
                    parsed_at=datetime.now().isoformat(timespec="seconds"),
                )
                self.db.upsert_document(
                    attachment_id, project["id"], parsed.title, parsed.text, parsed.page_count,
                    parsed.parser, identifiers, parsed.metadata,
                )
                result["parsed"] += 1
            except (OSError, ValueError, RuntimeError, requests.RequestException, SourceCircuitOpen) as exc:
                self.db.update_attachment(attachment_id, status="FAILED", error=str(exc)[:1000])
                result["failed"] += 1
        return result

    @staticmethod
    def discover_links(html: str, page_url: str) -> list[AttachmentLink]:
        soup = BeautifulSoup(html, "html.parser")
        links: dict[str, AttachmentLink] = {}
        for anchor in soup.select("a[href]"):
            href = (anchor.get("href") or "").strip()
            if not href or href.lower().startswith(("javascript:", "mailto:")):
                continue
            url = urljoin(page_url, href).split("#", 1)[0]
            text = (anchor.get("title") or anchor.get_text(" ", strip=True)).strip()
            path_name = Path(unquote(urlsplit(url).path)).name
            suffix = Path(path_name).suffix.lower()
            marked = any(marker in text for marker in DOCUMENT_LINK_MARKERS)
            download_path = any(marker in urlsplit(url).path.lower() for marker in ("attachment", "download", "upload"))
            if suffix not in SUPPORTED_EXTENSIONS and not (marked and download_path):
                continue
            file_name = path_name or text or f"attachment-{sha256(url.encode()).hexdigest()[:12]}"
            if Path(file_name).suffix.lower() not in SUPPORTED_EXTENSIONS:
                text_suffix = next((ext for ext in SUPPORTED_EXTENSIONS if ext in text.lower()), "")
                file_name += text_suffix
            links[url] = AttachmentLink(url, sanitize_file_name(file_name))
        return list(links.values())

    def download(self, link: AttachmentLink, attachment_id: int) -> tuple[Path, str, int, str]:
        suffix = Path(link.file_name).suffix.lower()
        destination = self.data_dir / f"{attachment_id}-{sanitize_file_name(link.file_name)}"
        part = destination.with_suffix(destination.suffix + ".part")
        sessions = [self.network.requests_session()]
        if self.network.has_direct_fallback:
            sessions.append(self.network.requests_session(direct=True))
        last_error: Exception | None = None
        for session in sessions:
            try:
                HostRateLimiter.wait(link.url, self.minimum_interval)
                with session.get(link.url, timeout=self.timeout, stream=True) as response:
                    if response.status_code in {403, 418, 429}:
                        raise SourceCircuitOpen(f"附件访问受限 HTTP {response.status_code}，已停止")
                    response.raise_for_status()
                    declared = int(response.headers.get("Content-Length") or 0)
                    if declared > self.max_bytes:
                        raise ValueError(f"附件超过大小限制：{declared} bytes")
                    digest, size = sha256(), 0
                    with part.open("wb") as stream:
                        for chunk in response.iter_content(1024 * 1024):
                            if not chunk:
                                continue
                            size += len(chunk)
                            if size > self.max_bytes:
                                raise ValueError("附件下载超过大小限制")
                            digest.update(chunk)
                            stream.write(chunk)
                    content_type = response.headers.get("Content-Type", "").split(";", 1)[0]
                part.replace(destination)
                if suffix not in SUPPORTED_EXTENSIONS:
                    inferred = mimetypes.guess_extension(content_type) or ""
                    if inferred not in SUPPORTED_EXTENSIONS:
                        raise ValueError(f"附件格式不受支持：{content_type or '未知'}")
                    renamed = destination.with_suffix(inferred)
                    destination.replace(renamed)
                    destination = renamed
                return destination, content_type, size, digest.hexdigest()
            except requests.exceptions.SSLError as exc:
                last_error = exc
                break
            except (requests.RequestException, OSError) as exc:
                last_error = exc
                continue
            finally:
                if part.exists():
                    part.unlink(missing_ok=True)
        if isinstance(last_error, requests.exceptions.SSLError):
            return self._download_with_windows_tls(link, attachment_id)
        raise RuntimeError(f"附件下载失败：{last_error}")

    def _download_with_windows_tls(self, link: AttachmentLink, attachment_id: int):
        destination = self.data_dir / f"{attachment_id}-{sanitize_file_name(link.file_name)}"
        command = [
            "curl.exe", "--fail", "--silent", "--show-error", "--location", "--ssl-no-revoke",
            "--max-time", str(self.timeout), "--max-filesize", str(self.max_bytes),
            "--output", str(destination), link.url,
        ]
        completed = subprocess.run(command, capture_output=True, timeout=self.timeout + 10)
        if completed.returncode:
            raise RuntimeError(completed.stderr.decode("utf-8", errors="replace").strip())
        size = destination.stat().st_size
        if size > self.max_bytes:
            destination.unlink(missing_ok=True)
            raise ValueError("附件下载超过大小限制")
        payload_hash = sha256(destination.read_bytes()).hexdigest()
        return destination, mimetypes.guess_type(destination.name)[0] or "", size, payload_hash

    @staticmethod
    def _host_allowed(page_url: str, attachment_url: str) -> bool:
        page_host = (urlsplit(page_url).hostname or "").lower()
        target_host = (urlsplit(attachment_url).hostname or "").lower()
        if urlsplit(attachment_url).scheme not in {"http", "https"}:
            return False
        if target_host == page_host:
            return True
        allowed = [item.strip().lower() for item in os.getenv("RADAR_DOCUMENT_ALLOWED_HOST_SUFFIXES", "").split(",") if item.strip()]
        if page_host.endswith(".beijing.gov.cn"):
            allowed.append(".beijing.gov.cn")
        return any(target_host == suffix.lstrip(".") or target_host.endswith(suffix) for suffix in allowed)


class DocumentRetentionManager:
    """Remove aged local binaries while preserving extracted text and FTS rows."""

    def __init__(self, db: Database, documents_dir: Path | None = None):
        self.db = db
        self.documents_dir = (documents_dir or DATA_DIR / "documents").resolve()

    def cleanup(self, retention_days: int = 2, dry_run: bool = False, now: datetime | None = None) -> dict:
        if retention_days < 0:
            raise ValueError("附件保留天数不能为负数")
        cutoff = (now or datetime.now()) - timedelta(days=retention_days)
        rows = self.db.attachments_for_cleanup(cutoff.isoformat(timespec="seconds"))
        result = {"eligible": len(rows), "deleted": 0, "missing": 0, "refused": 0, "failed": 0}
        for row in rows:
            path = Path(row["local_path"])
            try:
                resolved = path.resolve(strict=False)
                if not resolved.is_relative_to(self.documents_dir) or resolved == self.documents_dir:
                    result["refused"] += 1
                    continue
                if dry_run:
                    continue
                if resolved.exists():
                    if not resolved.is_file():
                        result["refused"] += 1
                        continue
                    resolved.unlink()
                    result["deleted"] += 1
                else:
                    result["missing"] += 1
                self.db.update_attachment(row["id"], status="PURGED", local_path="", error="")
            except OSError:
                result["failed"] += 1
        return result


def sanitize_file_name(value: str) -> str:
    clean = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", unquote(value or "attachment"))
    return clean.strip(" .")[:180] or "attachment"


def extract_identifiers(text: str) -> list[tuple[str, str]]:
    values: list[tuple[str, str]] = []
    for kind, pattern in IDENTIFIER_PATTERNS:
        for match in pattern.finditer(text or ""):
            value = match.group(1) if match.lastindex else match.group(0)
            value = value.strip("：:，,。;；()（）[]【】")
            if value:
                values.append((value, kind))
    return list(dict.fromkeys(values))


def fts_query(value: str) -> str:
    terms = re.findall(r"[\w\u4e00-\u9fff-]+", value or "", flags=re.UNICODE)
    return " AND ".join('"' + term.replace('"', '""') + '"' for term in terms if term)
